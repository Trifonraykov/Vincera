"""Tests for vincera.discovery — scanner, filesystem, database, spreadsheet, company model, agent."""

from __future__ import annotations

import asyncio
import csv
import os
import sqlite3
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

import pytest


# ============================================================
# Helpers
# ============================================================


def _run(coro):
    """Run an async coroutine synchronously."""
    return asyncio.run(coro)


def _mock_llm():
    """Create a mock OpenRouterClient."""
    llm = MagicMock()
    llm.think = AsyncMock(return_value="Here's a summary of your business.")
    llm.think_structured = AsyncMock(return_value={
        "business_type": "ecommerce",
        "industry": "retail",
        "confidence": 0.85,
        "software_stack": [{"name": "PostgreSQL", "category": "database", "role": "primary datastore"}],
        "data_architecture": [{"source": "orders.db", "type": "sqlite", "description": "Order data"}],
        "detected_processes": [{"name": "order processing", "manual": True, "frequency": "daily", "evidence": "spreadsheets"}],
        "automation_opportunities": [{"name": "invoice generation", "description": "Auto-generate invoices", "estimated_hours_saved": 5.0, "complexity": "medium", "evidence": "csv files"}],
        "pain_points": ["Manual data entry"],
        "risk_areas": ["Single point of failure on local DB"],
        "key_findings": ["Business runs on PostgreSQL + spreadsheets"],
    })
    return llm


def _mock_supabase():
    """Create a mock SupabaseManager."""
    sb = MagicMock()
    sb._company_id = "comp-123"
    sb.send_message.return_value = {"id": "msg-1"}
    sb.add_knowledge.return_value = {"id": "k-1"}
    sb.add_playbook_entry.return_value = {"id": "pb-1"}
    sb.query_playbook.return_value = []
    sb.query_knowledge.return_value = []
    sb.create_decision.return_value = "dec-1"
    sb.update_agent_status = MagicMock()
    sb.log_event.return_value = {"id": "ev-1"}
    return sb


def _mock_state(tmp_path: Path):
    """Create a mock GlobalState."""
    state = MagicMock()
    state.add_action = MagicMock()
    state.update_agent_status = MagicMock()
    state.get_agent_status.return_value = {"agent_name": "discovery", "status": "idle", "current_task": "none"}
    state._db = MagicMock()
    state._db.query.return_value = []
    return state


def _mock_settings(tmp_path: Path):
    """Create a mock VinceraSettings."""
    settings = MagicMock()
    settings.home_dir = tmp_path / "VinceraHQ"
    settings.home_dir.mkdir(parents=True, exist_ok=True)
    for subdir in ("agents", "knowledge", "core"):
        (settings.home_dir / subdir).mkdir(parents=True, exist_ok=True)
    settings.company_name = "TestCorp"
    settings.agent_name = "vincera"
    return settings


def _mock_verifier():
    """Create a mock Verifier."""
    from vincera.verification.verifier import CheckResult, VerificationResult

    v = MagicMock()
    v.verify = AsyncMock(return_value=VerificationResult(
        passed=True, checks=[CheckResult(name="test", passed=True, reason="ok")],
        confidence=0.95, blocked_reason=None,
    ))
    return v


def _mock_platform_service():
    """Create a mock PlatformService."""
    from vincera.platform import DiscoveryResult, ProcessInfo, ShareInfo, SoftwareInfo, TaskInfo

    svc = MagicMock()
    svc.list_installed_software.return_value = DiscoveryResult(
        items=[
            SoftwareInfo(name="python", version="3.11.0", source="brew"),
            SoftwareInfo(name="postgresql", version="16.0", source="brew"),
            SoftwareInfo(name="nginx", version="1.25", source="brew"),
        ],
        complete=True,
        errors=[],
    )
    svc.list_running_processes.return_value = DiscoveryResult(
        items=[
            ProcessInfo(pid=1, name="python3", user="user", cpu_percent=1.0, memory_percent=2.0, cmdline=["python3"]),
            ProcessInfo(pid=2, name="postgres", user="postgres", cpu_percent=0.5, memory_percent=3.0, cmdline=["postgres"]),
        ],
        complete=True,
        errors=[],
    )
    svc.list_network_shares.return_value = DiscoveryResult(items=[], complete=True, errors=[])
    svc.list_scheduled_tasks.return_value = DiscoveryResult(
        items=[TaskInfo(name="backup", schedule="daily", command="/usr/bin/backup.sh", status="running")],
        complete=True,
        errors=[],
    )
    return svc


# ============================================================
# Scanner tests
# ============================================================


class TestScannerEnvironment:
    def test_has_system_info(self) -> None:
        from vincera.discovery.scanner import SystemScanner

        svc = _mock_platform_service()
        scanner = SystemScanner(svc)
        env = _run(scanner.scan_environment())
        assert env.os_name is not None
        assert env.cpu_cores > 0
        assert env.ram_total_gb > 0

    def test_software_returns_list(self) -> None:
        from vincera.discovery.scanner import SystemScanner

        svc = _mock_platform_service()
        scanner = SystemScanner(svc)
        result = _run(scanner.scan_installed_software())
        assert len(result.data) > 0
        # Check enrichment: at least one has category
        categories = [s.get("category") for s in result.data if s.get("category")]
        assert len(categories) > 0

    def test_processes_returns_list(self) -> None:
        from vincera.discovery.scanner import SystemScanner

        svc = _mock_platform_service()
        scanner = SystemScanner(svc)
        result = _run(scanner.scan_running_processes())
        assert len(result.data) > 0
        # postgres should be tagged as database
        db_procs = [p for p in result.data if p.get("category") == "database"]
        assert len(db_procs) >= 1


# ============================================================
# Filesystem tests
# ============================================================


class TestFilesystem:
    def test_map_directory(self, tmp_path: Path) -> None:
        from vincera.discovery.filesystem import FilesystemMapper

        # Create known structure
        (tmp_path / "src").mkdir()
        (tmp_path / "src" / "main.py").write_text("pass")
        (tmp_path / "README.md").write_text("# readme")
        (tmp_path / "data").mkdir()
        (tmp_path / "data" / "report.csv").write_text("a,b,c")

        mapper = FilesystemMapper()
        tree = _run(mapper.map_directory(tmp_path, max_depth=2))
        assert tree.total_files >= 2
        assert tree.total_dirs >= 2

    def test_skips_hidden(self, tmp_path: Path) -> None:
        from vincera.discovery.filesystem import FilesystemMapper

        (tmp_path / ".hidden").mkdir()
        (tmp_path / ".hidden" / "secret.txt").write_text("shh")
        (tmp_path / "visible.txt").write_text("hello")

        mapper = FilesystemMapper()
        tree = _run(mapper.map_directory(tmp_path, max_depth=2))
        all_names = [e.name for e in tree.entries]
        assert ".hidden" not in all_names
        assert "visible.txt" in all_names

    def test_skips_node_modules(self, tmp_path: Path) -> None:
        from vincera.discovery.filesystem import FilesystemMapper

        (tmp_path / "node_modules").mkdir()
        (tmp_path / "node_modules" / "pkg").mkdir()
        (tmp_path / "app.js").write_text("console.log('hi')")

        mapper = FilesystemMapper()
        tree = _run(mapper.map_directory(tmp_path, max_depth=2))
        all_names = [e.name for e in tree.entries]
        assert "node_modules" not in all_names

    def test_never_reads_contents(self, tmp_path: Path) -> None:
        from vincera.discovery.filesystem import FilesystemMapper

        secret_file = tmp_path / "secret.txt"
        secret_file.write_text("TOP SECRET DATA")

        open_calls: list[str] = []
        original_open = open

        def tracking_open(path, *args, **kwargs):
            path_str = str(path)
            if str(tmp_path) in path_str:
                open_calls.append(path_str)
            return original_open(path, *args, **kwargs)

        mapper = FilesystemMapper()
        with patch("builtins.open", side_effect=tracking_open):
            _run(mapper.map_directory(tmp_path, max_depth=2))

        assert len(open_calls) == 0, f"open() was called on: {open_calls}"

    def test_handles_permission_error(self, tmp_path: Path) -> None:
        from vincera.discovery.filesystem import FilesystemMapper

        (tmp_path / "allowed.txt").write_text("ok")

        mapper = FilesystemMapper()
        with patch("os.scandir", side_effect=PermissionError("no access")):
            tree = _run(mapper.map_directory(tmp_path, max_depth=2))
        # Should return partial result, not crash
        assert tree is not None

    def test_identify_projects(self, tmp_path: Path) -> None:
        from vincera.discovery.filesystem import FilesystemMapper

        # Node project
        node_dir = tmp_path / "webapp"
        node_dir.mkdir()
        (node_dir / "package.json").write_text('{"name": "webapp"}')

        # Python project
        py_dir = tmp_path / "api"
        py_dir.mkdir()
        (py_dir / "requirements.txt").write_text("flask\n")

        mapper = FilesystemMapper()
        # Need to map first so we have trees to scan
        tree = _run(mapper.map_directory(tmp_path, max_depth=3))
        projects = _run(mapper.identify_project_structures(trees=[tree]))
        types = [p.project_type for p in projects]
        assert "node" in types or "javascript" in types
        assert "python" in types


# ============================================================
# Database tests
# ============================================================


class TestDatabase:
    def test_discover_from_processes(self) -> None:
        from vincera.discovery.database import DatabaseDiscovery
        from vincera.platform import ProcessInfo

        processes = [
            {"pid": 1, "name": "postgres", "cmdline": ["postgres", "-D", "/var/lib/pg"], "category": "database"},
        ]
        db_disc = DatabaseDiscovery()
        result = _run(db_disc.discover_databases(processes))
        assert len(result) >= 1
        assert result[0].db_type == "postgresql"

    def test_extract_schema_sqlite(self, tmp_path: Path) -> None:
        from vincera.discovery.database import DatabaseDiscovery, DatabaseInfo

        # Create test SQLite DB
        db_path = tmp_path / "test.db"
        conn = sqlite3.connect(str(db_path))
        conn.execute("CREATE TABLE orders (id INTEGER PRIMARY KEY, product TEXT, amount REAL)")
        conn.execute("CREATE TABLE customers (id INTEGER PRIMARY KEY, name TEXT, email TEXT)")
        conn.execute("INSERT INTO orders VALUES (1, 'Widget', 9.99)")
        conn.execute("INSERT INTO orders VALUES (2, 'Gadget', 19.99)")
        conn.execute("INSERT INTO customers VALUES (1, 'Alice', 'alice@test.com')")
        conn.commit()
        conn.close()

        db_info = DatabaseInfo(name="test.db", db_type="sqlite", port=None, version=None, path=str(db_path))
        db_disc = DatabaseDiscovery()
        schema = _run(db_disc.extract_schema(db_info))

        assert schema is not None
        assert len(schema.tables) == 2
        table_names = {t.name for t in schema.tables}
        assert "orders" in table_names
        assert "customers" in table_names

        orders_table = next(t for t in schema.tables if t.name == "orders")
        assert orders_table.row_count == 2
        col_names = [c.name for c in orders_table.columns]
        assert "id" in col_names
        assert "product" in col_names

    def test_schema_never_reads_data(self, tmp_path: Path) -> None:
        from vincera.discovery.database import DatabaseDiscovery, DatabaseInfo

        db_path = tmp_path / "secret.db"
        conn = sqlite3.connect(str(db_path))
        conn.execute("CREATE TABLE secrets (id INTEGER PRIMARY KEY, password TEXT)")
        conn.execute("INSERT INTO secrets VALUES (1, 'super_secret_password_123')")
        conn.commit()
        conn.close()

        db_info = DatabaseInfo(name="secret.db", db_type="sqlite", port=None, version=None, path=str(db_path))
        db_disc = DatabaseDiscovery()
        schema = _run(db_disc.extract_schema(db_info))

        assert schema is not None
        # Verify we got schema info
        assert len(schema.tables) == 1
        assert schema.tables[0].name == "secrets"
        # The schema should NOT contain actual data values
        schema_str = str(schema.model_dump())
        assert "super_secret_password_123" not in schema_str


# ============================================================
# Spreadsheet tests
# ============================================================


class TestSpreadsheet:
    def test_headers_csv(self, tmp_path: Path) -> None:
        from vincera.discovery.spreadsheet import SpreadsheetScanner

        csv_path = tmp_path / "orders.csv"
        with open(csv_path, "w", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(["Order ID", "Product", "Quantity", "Price"])
            writer.writerow([1, "Widget", 5, 9.99])
            writer.writerow([2, "Gadget", 3, 19.99])

        llm = _mock_llm()
        scanner = SpreadsheetScanner(llm)
        result = _run(scanner.scan_headers([csv_path]))

        assert len(result) == 1
        assert result[0].headers == ["Order ID", "Product", "Quantity", "Price"]
        assert result[0].estimated_row_count >= 2

    def test_headers_xlsx(self, tmp_path: Path) -> None:
        from vincera.discovery.spreadsheet import SpreadsheetScanner

        xlsx_path = tmp_path / "data.xlsx"
        try:
            import openpyxl

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Sales"
            ws.append(["Date", "Customer", "Amount", "Status"])
            ws.append(["2025-01-01", "Alice", 100, "paid"])
            ws.append(["2025-01-02", "Bob", 200, "pending"])
            wb.save(str(xlsx_path))

            llm = _mock_llm()
            scanner = SpreadsheetScanner(llm)
            result = _run(scanner.scan_headers([xlsx_path]))

            assert len(result) == 1
            assert result[0].headers == ["Date", "Customer", "Amount", "Status"]
            assert "Sales" in result[0].sheet_names
        except ImportError:
            pytest.skip("openpyxl not installed")

    def test_skips_bad_files(self, tmp_path: Path) -> None:
        from vincera.discovery.spreadsheet import SpreadsheetScanner

        bad_path = tmp_path / "corrupted.csv"
        bad_path.write_bytes(b"\x00\x01\x02\xff\xfe")

        llm = _mock_llm()
        scanner = SpreadsheetScanner(llm)
        result = _run(scanner.scan_headers([bad_path]))
        # Should not crash, just skip
        assert isinstance(result, list)


# ============================================================
# Company model tests
# ============================================================


class TestCompanyModel:
    def test_build(self) -> None:
        from vincera.discovery.company_model import CompanyModelBuilder

        llm = _mock_llm()
        builder = CompanyModelBuilder(llm)
        model = _run(builder.build(
            env=MagicMock(os_name="macOS"),
            software=MagicMock(data=[{"name": "postgres", "category": "database"}]),
            processes=MagicMock(data=[{"name": "postgres", "category": "database"}]),
            tasks=MagicMock(data=[]),
            filesystem=[],
            databases=[],
            spreadsheets=[],
        ))
        assert model.business_type == "ecommerce"
        assert model.confidence > 0
        assert len(model.automation_opportunities) > 0

    def test_narration(self) -> None:
        from vincera.discovery.company_model import CompanyModelBuilder

        llm = _mock_llm()
        builder = CompanyModelBuilder(llm)
        model = _run(builder.build(
            env=MagicMock(os_name="macOS"),
            software=MagicMock(data=[]),
            processes=MagicMock(data=[]),
            tasks=MagicMock(data=[]),
            filesystem=[],
            databases=[],
            spreadsheets=[],
        ))
        narration = _run(builder.to_narration(model))
        assert isinstance(narration, str)
        assert len(narration) > 0


# ============================================================
# Discovery Agent tests
# ============================================================


class TestDiscoveryAgent:
    def _make_agent(self, tmp_path: Path):
        from vincera.discovery.company_model import CompanyModel, CompanyModelBuilder
        from vincera.discovery.database import DatabaseDiscovery
        from vincera.discovery.filesystem import FilesystemMapper
        from vincera.discovery.network import NetworkDiscovery
        from vincera.discovery.scanner import EnvironmentInfo, ScanResult, SystemScanner
        from vincera.discovery.spreadsheet import SpreadsheetScanner
        from vincera.agents.discovery import DiscoveryAgent

        settings = _mock_settings(tmp_path)
        llm = _mock_llm()
        sb = _mock_supabase()
        state = _mock_state(tmp_path)
        verifier = _mock_verifier()

        # Mock scanner
        scanner = MagicMock(spec=SystemScanner)
        scanner.scan_environment = AsyncMock(return_value=EnvironmentInfo(
            os_name="macOS", os_version="14.0", hostname="test-host",
            cpu_model="Apple M1", cpu_cores=8, ram_total_gb=16.0, ram_available_gb=8.0,
            disk_partitions=[], docker_available=True, python_version="3.11.0",
            node_version="20.0.0", network_interfaces=[],
        ))
        scanner.scan_installed_software = AsyncMock(return_value=ScanResult(
            data=[
                {"name": "postgres", "version": "16.0", "source": "brew", "category": "database"},
                {"name": "python", "version": "3.11", "source": "brew", "category": "development"},
            ],
        ))
        scanner.scan_running_processes = AsyncMock(return_value=ScanResult(
            data=[
                {"pid": 1, "name": "postgres", "category": "database"},
                {"pid": 2, "name": "python3", "category": None},
            ],
        ))
        scanner.scan_scheduled_tasks = AsyncMock(return_value=ScanResult(data=[]))

        # Mock filesystem
        filesystem = MagicMock(spec=FilesystemMapper)
        filesystem.map_standard_paths = AsyncMock(return_value=[])
        filesystem.identify_project_structures = AsyncMock(return_value=[])
        filesystem.get_summary = MagicMock(return_value={
            "total_files": 150, "total_dirs": 20,
            "files_by_extension": {".py": 50, ".js": 30, ".csv": 10},
        })

        # Mock network
        network = MagicMock(spec=NetworkDiscovery)
        network.discover_shares = AsyncMock(return_value=[])

        # Mock database
        database = MagicMock(spec=DatabaseDiscovery)
        database.discover_databases = AsyncMock(return_value=[])

        # Mock spreadsheet
        spreadsheet = MagicMock(spec=SpreadsheetScanner)
        spreadsheet.scan_headers = AsyncMock(return_value=[])
        spreadsheet.analyze_patterns = AsyncMock(return_value={"summary": "No patterns"})

        # Mock model builder
        model_builder = MagicMock(spec=CompanyModelBuilder)
        mock_model = MagicMock()
        mock_model.business_type = "ecommerce"
        mock_model.automation_opportunities = [{"name": "invoice"}]
        mock_model.save_local = MagicMock()
        mock_model.save_to_supabase = MagicMock()
        model_builder.build = AsyncMock(return_value=mock_model)
        model_builder.to_narration = AsyncMock(return_value="Your business is an ecommerce operation.")

        agent = DiscoveryAgent(
            name="discovery",
            company_id="comp-123",
            config=settings,
            llm=llm,
            supabase=sb,
            state=state,
            verifier=verifier,
            scanner=scanner,
            filesystem=filesystem,
            network=network,
            database=database,
            spreadsheet=spreadsheet,
            model_builder=model_builder,
        )
        return agent, sb, mock_model

    def test_initial_sends_narrations(self, tmp_path: Path) -> None:
        agent, sb, _ = self._make_agent(tmp_path)
        result = _run(agent.run({"mode": "initial"}))

        # Count narration messages
        narration_calls = [
            c for c in sb.send_message.call_args_list
            if len(c[0]) >= 4 and c[0][3] == "discovery_narration"
        ]
        assert len(narration_calls) >= 8, f"Expected >= 8 narrations, got {len(narration_calls)}"

    def test_initial_saves_model(self, tmp_path: Path) -> None:
        agent, _, mock_model = self._make_agent(tmp_path)
        _run(agent.run({"mode": "initial"}))
        mock_model.save_local.assert_called_once()
        mock_model.save_to_supabase.assert_called_once()

    def test_records_playbook(self, tmp_path: Path) -> None:
        agent, sb, _ = self._make_agent(tmp_path)
        _run(agent.run({"mode": "initial"}))
        sb.add_playbook_entry.assert_called()
